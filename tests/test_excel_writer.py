import asyncio
from datetime import datetime, timezone, timedelta
import shutil
import os
from types import SimpleNamespace
from uuid import uuid4

import pytest


from app.integrations.onedrive import excel_writer as ew


class FakeClient:
    def __init__(self):
        self.upload_calls = []
        self.sample_file = None
        self.metadata = None
        self.download_exception = None

    async def get_file_metadata(self, path):
        return self.metadata

    async def download_file(self, remote_path, local_path):
        if self.download_exception:
            raise self.download_exception
        if self.sample_file:
            shutil.copy(self.sample_file, local_path)
            return
        # create an empty file so openpyxl will create new workbook
        open(local_path, "wb").close()

    async def upload_file(self, local_path, remote_path):
        # record and keep the uploaded file path for inspection
        self.upload_calls.append((local_path, remote_path))


@pytest.mark.asyncio
async def test_upsert_creates_and_uploads(monkeypatch, tmp_path):
    fake = FakeClient()
    # no metadata, download will create new
    monkeypatch.setattr(ew, "OneDriveClient", lambda: fake)

    preventivo = SimpleNamespace(id=uuid4(), status="NEW", quote_data={"descrizione_lavori": "test"}, created_at=datetime.now(timezone.utc))
    customer = SimpleNamespace(name="Luca", email="luca@example.com", phone="3331234567")
    tenant = SimpleNamespace(id=uuid4())

    await ew.upsert_preventivo_row(preventivo, customer, tenant)

    assert len(fake.upload_calls) == 1
    uploaded_path = fake.upload_calls[0][0]
    assert os.path.exists(uploaded_path)

    # verify workbook has Preventivi sheet and header
    from openpyxl import load_workbook

    wb = load_workbook(uploaded_path)
    assert "Preventivi" in wb.sheetnames
    ws = wb["Preventivi"]
    header = [c.value for c in ws[1]]
    assert "id" in header and "status" in header


@pytest.mark.asyncio
async def test_upsert_aborts_if_recently_modified(monkeypatch):
    fake = FakeClient()
    now = datetime.now(timezone.utc)
    fake.metadata = {"lastModifiedDateTime": now.isoformat().replace("+00:00", "Z")}
    monkeypatch.setattr(ew, "OneDriveClient", lambda: fake)

    preventivo = SimpleNamespace(id=uuid4(), status="NEW", quote_data={}, created_at=datetime.now(timezone.utc))
    customer = SimpleNamespace(name="Test", email=None, phone=None)
    tenant = SimpleNamespace(id=uuid4())

    await ew.upsert_preventivo_row(preventivo, customer, tenant)

    # since file modified very recently, upload should not be called
    assert len(fake.upload_calls) == 0


@pytest.mark.asyncio
async def test_upsert_updates_existing_row(monkeypatch, tmp_path):
    # prepare a sample workbook with a Preventivi sheet and a row
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Preventivi"
    ws.append(["id", "tenant_id", "customer_name", "customer_email", "customer_phone", "descrizione_lavori", "status", "created_at", "updated_at"])

    pid = str(uuid4())
    ws.append([pid, "t1", "Old", "old@example.com", "", "old job", "OLD", datetime.now(timezone.utc).isoformat(), None])

    sample = tmp_path / "sample.xlsx"
    wb.save(sample)

    fake = FakeClient()
    fake.sample_file = str(sample)
    # set metadata old enough to allow update
    fake.metadata = {"lastModifiedDateTime": (datetime.now(timezone.utc) - timedelta(seconds=60)).isoformat().replace("+00:00", "Z")}
    monkeypatch.setattr(ew, "OneDriveClient", lambda: fake)

    # create preventivo matching pid with new status
    preventivo = SimpleNamespace(id=pid, status="SENT", quote_data={"descrizione_lavori": "updated job"}, created_at=datetime.now(timezone.utc))
    customer = SimpleNamespace(name="NewName", email="new@example.com", phone="333")
    tenant = SimpleNamespace(id=uuid4())

    await ew.upsert_preventivo_row(preventivo, customer, tenant)

    # upload called once and we can inspect uploaded file
    assert len(fake.upload_calls) == 1
    uploaded_path = fake.upload_calls[0][0]
    wb2 = load_workbook(uploaded_path)
    ws2 = wb2["Preventivi"]
    # find row with pid and check status updated
    found = False
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == pid:
            found = True
            assert row[6] == "SENT"
    assert found
